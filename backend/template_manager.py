from docx import Document

def populate_template(template_path, placeholders):
    """
    Populates a Word template with user input data.

    Args:
        template_path (str): Path to the .docx template.
        placeholders (dict): A dictionary with placeholders and user-provided values.

    Returns:
        doc: The populated Word document.
    """
    doc = Document(template_path)
    for p in doc.paragraphs:
        for key, value in placeholders.items():
            if key in p.text:
                p.text = p.text.replace(key, value)
    return doc
